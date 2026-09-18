#!/usr/bin/env python3
"""Convert Patterns_findings.xlsx into the data file the website reads.

Run after every change to the workbook:

    python tools/build_patterns.py

Reads   : Patterns_findings.xlsx      (the master; the team keeps editing it in Excel)
          *.ipynb                     (the notebooks, rendered for in-page reading)
Writes  : docs/data/patterns.js       (window.PATTERNS = {...}, loadable via file://)
          docs/data/notebooks.js      (window.NOTEBOOKS = {...}, the rendered notebooks)
          docs/data/build-report.txt  (data issues found, nothing is silently fixed)

Notebook links are resolved by name, not from the workbook: a pattern called
"Transaction Detection" links to transaction_detection.ipynb if that file exists
in the repository root. Committing a notebook is enough to make its link appear.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import unicodedata
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.rich_text import CellRichText

REPO = Path(__file__).resolve().parent.parent
WORKBOOK = REPO / "Patterns_findings.xlsx"
SHEET = "Pattern Definitions & Examples "
# Per-event-log evidence and its analyst utility live on their own sheets, both
# joined to the catalogue by No. A pattern x log block is shown only when the
# Utility sheet has content for it; the evidence is then appended with its utility.
EVIDENCE_SHEET = "EventLog Analysis Results"
UTILITY_SHEET = "Utility"
OUT_JS = REPO / "docs" / "data" / "patterns.js"
OUT_NOTEBOOKS = REPO / "docs" / "data" / "notebooks.js"
OUT_REPORT = REPO / "docs" / "data" / "build-report.txt"

HEADER_ROW = 9
FIRST_DATA_ROW = 11

# Website field key -> exact header text in the workbook's header row (row 9).
# Columns are matched by header, not by position, so they can be reordered in
# Excel without touching this file. Matching is case-insensitive and ignores
# differences in surrounding/internal whitespace. Any column whose header is not
# listed here (Relevant Evidence, Literature Support, Notes, …) is deliberately
# not carried onto the site: the evidence comes from EVIDENCE_SHEET instead.
FIELD_HEADERS = {
    "number": "No.",
    "category": "Category",
    "tags": "Tags",
    "name": "Name",
    "information_need": "Information Need / Modification Objective",
    "motivation": "Motivation / Context",
    "preconditions": "Preconditions",
    "approach": "Approach",
    "output": "Output",
    "dependencies": "Dependencies",
    "imperfection_relation": "Relation with Imperfection Patterns",
}

# Fields whose cells carry formatting worth keeping (bold/italic, ink colours,
# line breaks); the rest are read as plain text.
RICH_FIELDS = {
    "information_need", "motivation", "preconditions", "approach", "output",
    "imperfection_relation",
}

report_lines: list[str] = []


def note(kind: str, message: str) -> None:
    report_lines.append(f"{kind:<9} {message}")


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.replace("–", " ").replace("—", " ").replace("&", " and ")
    text = re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_").lower()
    return text


def esc(text: str) -> str:
    return (
        text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )


def rgb_of(colour) -> str | None:
    """Return an RRGGBB string for an explicit rgb colour, else None."""
    if colour is None:
        return None
    try:
        value = colour.rgb
    except (AttributeError, TypeError):
        return None
    if isinstance(value, str) and len(value) == 8:
        return value[2:].upper()
    return None


def rich_to_html(value) -> str:
    """Render a cell as HTML, preserving bold/italic/underline and red/blue runs.

    Red text is Lisa's inline remarks, blue is Marco's (legend in H3:H4), so the
    colour carries meaning and is kept as a marked-up span rather than dropped.
    """
    if value is None:
        return ""
    if not isinstance(value, CellRichText):
        return esc(str(value)).replace("\n", "<br>")

    out = []
    for block in value:
        text = esc(str(block))
        font = getattr(block, "font", None)
        if font is not None:
            if font.b:
                text = f"<strong>{text}</strong>"
            if font.i:
                text = f"<em>{text}</em>"
            if font.u:
                text = f"<u>{text}</u>"
            colour = rgb_of(getattr(font, "color", None))
            if colour == "FF0000":
                text = f'<span class="ink ink-red">{text}</span>'
            elif colour in {"0070C0", "0000FF", "1F4E79", "2E75B6", "0563C1"}:
                text = f'<span class="ink ink-blue">{text}</span>'
        out.append(text)
    return "".join(out).replace("\n", "<br>")


def plain(value) -> str:
    return "" if value is None else str(value).strip()


# Evidence cells holding one of these (after stripping tags, whitespace and a
# trailing full stop, case-insensitively) carry no real content and are dropped.
BLANK_EVIDENCE = {"", "n/a", "na", "nothing relevant identified"}

# Columns on the evidence sheet that are not event logs.
EVIDENCE_META_HEADERS = {"no.", "no", "number", "category", "pattern name", "name"}


def _tidy(html: str) -> str:
    html = re.sub(r"^(?:\s|<br>)+|(?:\s|<br>)+$", "", html)
    return html.strip()


def _is_blank_evidence(html: str) -> bool:
    """True when an evidence cell says nothing worth showing."""
    text = re.sub(r"<[^>]+>", " ", html or "")
    text = re.sub(r"\s+", " ", text).strip().rstrip(".").strip().lower()
    return text in BLANK_EVIDENCE


def _pretty_dataset(header: str) -> str:
    """Tidy an event-log column header: one space after 'BPIC', whitespace collapsed."""
    label = re.sub(r"\s+", " ", str(header or "")).strip()
    return re.sub(r"^(BPIC)\s*(?=\d)", r"\1 ", label)


def read_log_matrix(wb, sheet_name: str) -> tuple[dict[str, dict[str, str]], dict[str, str], list[str]]:
    """Read a per-pattern x per-event-log sheet, keyed by pattern No.

    Both "EventLog Analysis Results" (evidence) and "Utility" share this layout:
    a No./Category/Pattern name block followed by one column per event log.
    Returns (by_number, name_by_number, log_order) where by_number[number] maps an
    event-log label to that cell's HTML for every column that holds real content
    (empty, "N/A" and "Nothing relevant identified" cells are dropped), and
    log_order is the event-log labels in column order. Columns are located by
    header, so the sheet can be reordered, and the join to the catalogue is on the
    pattern number.
    """
    if sheet_name not in wb.sheetnames:
        note("evidence", f"sheet {sheet_name!r} not found; nothing loaded from it")
        return {}, {}, []
    ws = wb[sheet_name]

    num_col = name_col = None
    log_cols: list[tuple[int, str]] = []
    for col in range(1, ws.max_column + 1):
        header = re.sub(r"\s+", " ", str(ws.cell(1, col).value or "")).strip()
        key = header.lower()
        if key in {"no.", "no", "number"}:
            num_col = col
        elif key in {"pattern name", "name"}:
            name_col = col
        elif key and key not in EVIDENCE_META_HEADERS:
            log_cols.append((col, _pretty_dataset(header)))
    if num_col is None:
        note("evidence", f"{sheet_name}: no 'No.' column found; nothing loaded from it")
        return {}, {}, []

    by_number: dict[str, dict[str, str]] = {}
    names: dict[str, str] = {}
    for row in range(2, ws.max_row + 1):
        number = plain(ws.cell(row, num_col).value)
        # Skip section-header rows: a No. (e.g. "1") with a category but no
        # pattern name. Real rows always name a pattern.
        row_name = plain(ws.cell(row, name_col).value) if name_col is not None else ""
        if not number or (name_col is not None and not row_name):
            continue
        if number in by_number:
            note("evidence", f"{sheet_name}: pattern number {number} appears in more than one row")
        cells = {}
        for col, label in log_cols:
            html = _tidy(rich_to_html(ws.cell(row, col).value))
            if not _is_blank_evidence(html):
                cells[label] = html
        by_number[number] = cells
        if row_name:
            names[number] = row_name
    return by_number, names, [label for _, label in log_cols]


def split_tags(raw: str) -> list[str]:
    """Tags are free text: one per line and/or comma-separated, order preserved.

    A value like "Analysis\\nModification" becomes two tags. Duplicates within a
    cell are collapsed while keeping the first occurrence's order.
    """
    tags: list[str] = []
    for line in (raw or "").split("\n"):
        for piece in line.split(","):
            piece = piece.strip()
            if piece and piece not in tags:
                tags.append(piece)
    return tags


def split_dependencies(raw: str) -> list[str]:
    """Dependencies are free text: one per line, sometimes comma-separated.

    Deliberately not split on a bare " and ", because pattern names contain it
    ("Log and Case Attribute Candidates").
    """
    if not raw:
        return []
    parts: list[str] = []
    for line in raw.split("\n"):
        for piece in line.split(","):
            piece = re.sub(r"^\s*and\b\s*", "", piece.strip()).strip().rstrip(".")
            if piece:
                parts.append(piece)
    return parts


def workbook_saved(path: Path) -> str:
    """When the workbook was last saved, not when the build ran.

    Deriving the stamp from the input keeps the build reproducible: rerunning it
    on unchanged input produces byte-identical output, so the GitHub Action does
    not commit a new timestamp on every run. It also answers the question the
    reader actually has - how current is this catalogue.
    """
    try:
        with zipfile.ZipFile(path) as zf:
            core = ET.fromstring(zf.read("docProps/core.xml"))
        for tag in ("modified", "created"):
            found = core.find(f"{{http://purl.org/dc/terms/}}{tag}")
            if found is not None and found.text:
                return found.text[:10]
    except (KeyError, ET.ParseError, zipfile.BadZipFile):
        pass
    return dt.date.fromtimestamp(path.stat().st_mtime).isoformat()


def _norm_header(value) -> str:
    """Normalise a header cell for matching: collapse whitespace, lower-case."""
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def resolve_columns(ws) -> dict[int, str]:
    """Map each website field to its workbook column by matching the header row.

    Returns {column index: field key}. A header listed in FIELD_HEADERS but not
    found in the sheet is reported (so a renamed or removed column is caught)
    rather than silently dropped.
    """
    header_to_col: dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        header = _norm_header(ws.cell(HEADER_ROW, col).value)
        if header:
            header_to_col.setdefault(header, col)

    columns: dict[int, str] = {}
    for key, header in FIELD_HEADERS.items():
        col = header_to_col.get(_norm_header(header))
        if col is None:
            note("column", f"header {header!r} not found in row {HEADER_ROW}; {key} left blank")
            continue
        columns[col] = key
    return columns


def build(workbook_path: Path) -> dict:
    wb = openpyxl.load_workbook(workbook_path, rich_text=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found; sheets are {wb.sheetnames}")
    ws = wb[SHEET]
    columns = resolve_columns(ws)
    evidence_by_number, evidence_names, evidence_logs = read_log_matrix(wb, EVIDENCE_SHEET)
    utility_by_number, utility_names, utility_logs = read_log_matrix(wb, UTILITY_SHEET)
    # Present the event logs in the evidence sheet's column order, with any that
    # only appear on the Utility sheet appended after them.
    log_order = evidence_logs + [log for log in utility_logs if log not in evidence_logs]

    notebooks = {slugify(p.stem): p.name for p in sorted(REPO.glob("*.ipynb"))}
    used_notebooks: set[str] = set()

    patterns = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        if all(ws.cell(row, col).value in (None, "") for col in columns):
            continue

        # Every field defaults to blank, so a column missing from the sheet
        # yields an empty value instead of a KeyError downstream.
        record = {key: "" for key in FIELD_HEADERS}
        for col, key in columns.items():
            cell = ws.cell(row, col)
            if key in RICH_FIELDS:
                record[key] = rich_to_html(cell.value)
            else:
                record[key] = plain(cell.value)

        name = record["name"]
        if not name:
            note("skip", f"row {row}: no name, skipped")
            continue

        slug = slugify(name)
        record["slug"] = slug
        record["row"] = row
        record["tags"] = split_tags(record["tags"])
        # One block per event log, shown only where the Utility sheet has content;
        # each block carries the evidence with its utility appended. A log whose
        # utility cell is empty is omitted entirely, even if evidence exists.
        ev_cells = evidence_by_number.get(record["number"], {})
        ut_cells = utility_by_number.get(record["number"], {})
        record["evidence_blocks"] = [
            {"dataset": log, "html": ev_cells.get(log, ""), "utility": ut_cells[log]}
            for log in log_order
            if log in ut_cells
        ]
        record["datasets"] = sorted(
            {b["dataset"] for b in record["evidence_blocks"] if b["dataset"]}
        )
        record["dependency_names"] = split_dependencies(record["dependencies"])
        notebook = notebooks.get(slug)
        if notebook:
            used_notebooks.add(notebook)
        record["notebook"] = notebook

        patterns.append(record)

    # Resolve dependency names against the catalogue by pattern name.
    index = {slugify(p["name"]): p["slug"] for p in patterns}
    for p in patterns:
        resolved = []
        for dep in p["dependency_names"]:
            target = index.get(slugify(dep))
            if target and target != p["slug"]:
                resolved.append({"label": dep, "slug": target})
            else:
                resolved.append({"label": dep, "slug": None})
                if len(dep.split()) <= 8:
                    note("dependency", f"{p['name']}: dependency {dep!r} matches no pattern")
        p["dependency_links"] = resolved

    seen_numbers: dict[str, str] = {}
    for p in patterns:
        if p["number"] in seen_numbers:
            note("number", f"number {p['number']} used by both {seen_numbers[p['number']]!r} and {p['name']!r}")
        seen_numbers[p["number"]] = p["name"]

    for name in sorted(set(notebooks.values()) - used_notebooks):
        note("notebook", f"{name} is in the repo but matches no pattern name")

    # Cross-check the evidence and utility sheets against the catalogue (join is
    # on the number). A pattern with no utility shows no Evidence & Utility section.
    pattern_numbers = {p["number"] for p in patterns}
    for p in patterns:
        if p["number"] not in evidence_by_number:
            note("evidence", f"{p['name']} (No. {p['number']}): no row in {EVIDENCE_SHEET}")
        if not p["evidence_blocks"]:
            note("utility", f"{p['name']} (No. {p['number']}): no utility filled in, so no evidence is shown")
        # Utility written for a log that has no evidence to append it to.
        ev_cells = evidence_by_number.get(p["number"], {})
        for block in p["evidence_blocks"]:
            if not block["html"]:
                note("utility", f"{p['name']} (No. {p['number']}, {block['dataset']}): "
                                f"utility present but no evidence in {EVIDENCE_SHEET}")
        other = evidence_names.get(p["number"]) or utility_names.get(p["number"]) or ""
        if other and slugify(other) != slugify(p["name"]):
            note("evidence", f"No. {p['number']}: name differs between sheets "
                             f"({p['name']!r} vs {other!r})")
    for number in sorted((set(evidence_by_number) | set(utility_by_number)) - pattern_numbers):
        label = evidence_names.get(number) or utility_names.get(number) or ""
        note("evidence", f"row No. {number} ({label!r}) matches no pattern")

    # Order categories by the lowest pattern number they contain, so the
    # catalogue reads 1.x, 2.x, 3.x … rather than alphabetically. Categories
    # with no numeric pattern fall to the end; ties break on the name.
    def _num_key(value: str) -> float:
        try:
            return float(value)
        except (TypeError, ValueError):
            return float("inf")

    def _category_key(cat: str) -> tuple[float, str]:
        nums = [_num_key(p["number"]) for p in patterns if p["category"] == cat]
        return (min(nums, default=float("inf")), cat)

    categories = sorted({p["category"] for p in patterns if p["category"]}, key=_category_key)
    for cat in categories:
        members = [p["name"] for p in patterns if p["category"] == cat]
        if len(members) == 1:
            note("category", f"category {cat!r} has a single member: {members[0]}")

    return {
        "generated": workbook_saved(workbook_path),
        "source": workbook_path.name,
        "categories": categories,
        "tags": sorted({t for p in patterns for t in p["tags"]}),
        "datasets": sorted({d for p in patterns for d in p["datasets"]}),
        "patterns": patterns,
    }


# Output types that only exist to drive JavaScript in a live Jupyter session
# (tqdm progress bars, the VS Code data viewer). Dropping them before conversion
# makes nbconvert fall back to the plain-text output instead of emitting an
# empty box plus a few hundred kilobytes of widget state.
DEAD_MIMETYPES = (
    "application/vnd.jupyter.widget-view+json",
    "application/vnd.jupyter.widget-state+json",
    "application/vnd.microsoft.datawrangler.viewer.v0+json",
)


# nbconvert leaves TeX for a browser-side MathJax that this site deliberately
# does not load (it would need a CDN, and the site has to work offline). Instead
# the formulas are turned into readable plain text at build time. This is not a
# TeX engine and is not meant to be one: it covers the notation the notebooks
# actually use, and anything it does not know simply keeps its backslash.
MATH_BLOCK = re.compile(r"\$\$(.+?)\$\$", re.S)
# Inline math, kept on one line and never crossing a tag boundary.
MATH_INLINE = re.compile(r"(?<![\w$])\$(?!\s)([^$\n<>]{1,200}?)(?<!\s)\$(?![\w$])")

TEX_WORDS = {
    r"\\cdot": "·", r"\\times": "×", r"\\leq": "≤", r"\\geq": "≥",
    r"\\neq": "≠", r"\\approx": "≈", r"\\in\b": "∈", r"\\cup": "∪",
    r"\\cap": "∩", r"\\sum": "sum", r"\\prod": "prod", r"\\min": "min",
    r"\\max": "max", r"\\log": "log", r"\\left": "", r"\\right": "",
    r"\\,": "", r"\\;": " ", r"\\!": "", r"\\quad": "  ", r"\\qquad": "   ",
}


def _wrap(part: str) -> str:
    """Parenthesise a fraction operand only when it is more than a single term."""
    part = part.strip()
    return part if re.fullmatch(r"[\w.^_]+", part) else f"({part})"


def tex_to_text(tex: str) -> str:
    """Best-effort plain-text rendering of the TeX used in these notebooks."""
    out = tex.strip()
    out = re.sub(r"\\(?:mathrm|mathit|text|mathbf)\{([^{}]*)\}", r"\1", out)
    out = re.sub(r"\\frac\s*\{([^{}]*)\}\s*\{([^{}]*)\}",
                 lambda m: f"{_wrap(m.group(1))} / {_wrap(m.group(2))}", out)
    for pattern, replacement in TEX_WORDS.items():
        out = re.sub(pattern, replacement, out)
    out = re.sub(r"\s+", " ", out)
    return out.strip()


def mark_math(html: str) -> str:
    # Inline "$...$" is only interpreted in notebooks that also use display
    # math, so a stray pair of dollar amounts elsewhere is left alone.
    uses_tex = "$$" in html
    html = MATH_BLOCK.sub(
        lambda m: f'<code class="tex">{tex_to_text(m.group(1))}</code>', html)
    if uses_tex:
        html = MATH_INLINE.sub(
            lambda m: f'<code class="tex inline">{tex_to_text(m.group(1))}</code>', html)
    return html


def render_notebooks(patterns: list[dict]) -> dict[str, str]:
    """Render each linked notebook to an HTML fragment for in-page reading.

    Optional: without nbconvert the site simply keeps linking to the notebooks
    instead of showing them, so the build never fails over a missing dependency.
    """
    wanted = [p for p in patterns if p["notebook"]]
    if not wanted:
        return {}

    try:
        import nbformat
        from nbconvert import HTMLExporter
    except ImportError:
        note("notebook", "nbconvert is not installed - notebooks are linked but not shown "
                         "(pip install -r tools/requirements.txt)")
        return {}

    exporter = HTMLExporter(template_name="basic", embed_images=True)
    rendered = {}
    for p in wanted:
        path = REPO / p["notebook"]
        try:
            nb = nbformat.read(path, as_version=4)
        except Exception as exc:  # a malformed notebook must not break the build
            note("notebook", f"{p['notebook']}: could not be read ({exc})")
            continue

        for cell in nb.cells:
            for output in cell.get("outputs", []):
                data = output.get("data")
                if data:
                    for mime in DEAD_MIMETYPES:
                        data.pop(mime, None)

        try:
            body, _ = exporter.from_notebook_node(nb)
        except Exception as exc:
            note("notebook", f"{p['notebook']}: could not be rendered ({exc})")
            continue

        # The fragment is injected with innerHTML, which never runs scripts;
        # removing them keeps the markup honest rather than merely inert.
        body = re.sub(r"<script\b.*?</script>", "", body, flags=re.S | re.I)
        body = mark_math(body)
        rendered[p["slug"]] = body.strip()

    return rendered


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook", nargs="?", default=str(WORKBOOK),
                    help="path to the xlsx (default: Patterns_findings.xlsx in the repo root)")
    args = ap.parse_args()

    path = Path(args.workbook)
    if not path.exists():
        raise SystemExit(f"workbook not found: {path}")

    data = build(path)
    OUT_JS.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(data, ensure_ascii=False, indent=1)
    OUT_JS.write_text(
        "// Generated by tools/build_patterns.py - do not edit by hand.\n"
        "// Edit Patterns_findings.xlsx and re-run the script instead.\n"
        f"window.PATTERNS = {payload};\n",
        encoding="utf-8",
    )

    notebooks = render_notebooks(data["patterns"])
    OUT_NOTEBOOKS.write_text(
        "// Generated by tools/build_patterns.py - do not edit by hand.\n"
        "// Rendered from the .ipynb files in the repository root.\n"
        f"window.NOTEBOOKS = {json.dumps(notebooks, ensure_ascii=False)};\n",
        encoding="utf-8",
    )

    header = [
        f"Build report for {data['source']}, saved {data['generated']}",
        f"{len(data['patterns'])} patterns, "
        f"{sum(1 for p in data['patterns'] if p['notebook'])} with a notebook, "
        f"{len(notebooks)} notebooks rendered",
        "",
    ]
    body = report_lines or ["No issues found."]
    OUT_REPORT.write_text("\n".join(header + sorted(body)) + "\n", encoding="utf-8")

    print("\n".join(header + sorted(body)))
    print(f"\nwrote {OUT_JS.relative_to(REPO)}, {OUT_NOTEBOOKS.relative_to(REPO)} "
          f"and {OUT_REPORT.relative_to(REPO)}")


if __name__ == "__main__":
    main()
